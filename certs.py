"""Generate personalized PDF certificates; email them only with --send."""

import argparse
import csv
from datetime import date
from email.errors import HeaderParseError
from email.headerregistry import Address
from email.message import EmailMessage
from html import escape
import os
from pathlib import Path
import re
import smtplib
import ssl
from string import Formatter
import sys
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from defusedxml.common import DefusedXmlException
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent


def valid_email(value):
    if not value or not value.isascii() or any(char.isspace() or ord(char) < 32 or ord(char) == 127 for char in value):
        return False
    try:
        address = Address(addr_spec=value)
        return bool(address.username and address.domain)
    except (ValueError, IndexError, HeaderParseError):
        return False


def read_recipients(path, require_email=False):
    """Validate the entire input before creating files or contacting SMTP."""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as source:
            rows = list(csv.reader(source, strict=True))
    elif path.suffix.lower() == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = list(book.active.values)
        finally:
            book.close()
    else:
        raise ValueError("Input must be an .xlsx or UTF-8 .csv file.")

    headers = [str(value or "").strip().lower() for value in rows[0]] if rows else []
    for column in ("name", "email"):
        if headers.count(column) > 1:
            raise ValueError(f"Duplicate '{column}' column in input.")
    if "name" not in headers or (require_email and "email" not in headers):
        raise ValueError("Input needs a Name column, and an Email column when using --send.")

    recipients = []
    for number, row in enumerate(rows[1:], start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        values = dict(zip(headers, row))
        name = str(values.get("name") or "").strip()
        email = str(values.get("email") or "").strip()
        if not name or any(ord(char) < 32 or ord(char) == 127 for char in name):
            raise ValueError(f"Row {number}: Name must contain text without control characters.")
        if (require_email and not email) or (email and not valid_email(email)):
            raise ValueError(f"Row {number}: provide a valid Email address.")
        recipients.append((name, email))
    if not recipients:
        raise ValueError("Input contains no recipients.")
    return recipients


def render_certificate(template, name, certificate_id, args):
    image = template.copy()
    draw = ImageDraw.Draw(image)
    width, height = image.size
    position = args.name_position or (width / 2, height / 2)
    name_width = args.name_width or width * 0.8
    font = ImageFont.truetype(str(args.name_font), args.font_size)
    # Fit long names without changing their spelling or capitalization.
    while font.size > 1:
        left, top, right, bottom = draw.textbbox(position, name, font=font, anchor="mm")
        if right - left <= name_width:
            break
        font = font.font_variant(size=max(1, min(font.size - 1, int(font.size * name_width / (right - left)))))
    detail_font = ImageFont.truetype(str(args.detail_font), args.detail_size or max(10, round(width / 60)))
    labels = (
        ("Name", name, position, font, "mm"),
        ("Certificate ID", certificate_id, args.id_position or (width * 0.07, height * 0.9), detail_font, "lt"),
        ("Date", args.date, args.date_position or (width * 0.7, height * 0.9), detail_font, "lt"),
    )
    for label, text, position, font, anchor in labels:
        left, top, right, bottom = draw.textbbox(position, text, font=font, anchor=anchor)
        if left < 0 or top < 0 or right > width or bottom > height:
            raise ValueError(f"{label} for {name!r} falls outside the template; adjust its position or font size.")
        if label == "Name" and right - left > name_width:
            raise ValueError(f"Name {name!r} cannot fit within --name-width.")
        draw.text(position, text, fill=args.color, font=font, anchor=anchor)
    return image


def send_certificates(records, args):
    sender = os.environ.get("SENDER_EMAIL", "")
    password = os.environ.get("SENDER_EMAIL_PASSWORD", "")
    failures = 0
    with (args.output / "delivery.csv").open("x", encoding="utf-8", newline="") as log:
        writer = csv.writer(log)
        writer.writerow(("Certificate ID", "Status"))
        log.flush()
        with smtplib.SMTP_SSL(args.smtp_host, args.smtp_port, context=ssl.create_default_context(), timeout=30) as smtp:
            smtp.login(sender, password)
            for name, email, certificate_id, filename in records:
                message = EmailMessage()
                message["From"] = sender
                message["To"] = email
                message["Subject"] = args.subject
                body = args.message.format(name=name, certificate_id=certificate_id)
                message.set_content(body)
                if args.logo:
                    with Image.open(args.logo) as logo:
                        subtype = logo.format.lower()
                    message.add_alternative(f'<pre>{escape(body)}</pre><img src="cid:logo" alt="Organization logo">', subtype="html")
                    message.get_payload()[-1].add_related(args.logo.read_bytes(), maintype="image", subtype=subtype, cid="<logo>")
                message.add_attachment((args.output / filename).read_bytes(), maintype="application", subtype="pdf", filename=filename)
                try:
                    smtp.send_message(message)
                except (smtplib.SMTPException, OSError) as error:
                    failures += 1
                    status = "Failed or uncertain; check provider before retrying"
                    print(f"Email failed for {certificate_id}: {error}", file=sys.stderr)
                else:
                    status = "Accepted by SMTP server"
                writer.writerow((certificate_id, status))
                log.flush()
    return failures


def build_parser():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    parser.add_argument("--input", type=Path, default=Path("file.xlsx"), help="Recipient spreadsheet (.xlsx or .csv)")
    parser.add_argument("--template", type=Path, default=Path("cert.jpg"), help="Certificate background image")
    parser.add_argument("--output", type=Path, default=Path("out"), help="New or empty output directory")
    parser.add_argument("--prefix", default="CERT-", help="Certificate ID prefix (letters, digits, underscores, hyphens)")
    parser.add_argument("--start", type=int, default=1, help="First certificate number")
    parser.add_argument("--date", default=date.today().strftime("%d-%m-%Y"), help="Date text printed on certificates")
    parser.add_argument("--name-font", type=Path, default=ROOT / "fontname.ttf", help="Name font")
    parser.add_argument("--detail-font", type=Path, default=ROOT / "fontid.ttf", help="ID and date font")
    parser.add_argument("--font-size", type=int, default=250, help="Maximum name font size; shrinks to fit")
    parser.add_argument("--detail-size", type=int, help="ID and date font size (default: scaled to template width)")
    parser.add_argument("--name-width", type=int, help="Maximum name width in pixels (default: 80%% of image width)")
    parser.add_argument("--name-position", type=int, nargs=2, metavar=("X", "Y"), help="Center of the name in pixels")
    parser.add_argument("--id-position", type=int, nargs=2, metavar=("X", "Y"), help="Top-left of the ID in pixels")
    parser.add_argument("--date-position", type=int, nargs=2, metavar=("X", "Y"), help="Top-left of the date in pixels")
    parser.add_argument("--color", default="#0089d1", help="Text color")
    parser.add_argument("--dpi", type=int, default=300, help="PDF print resolution")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--preview", type=Path, metavar="PNG", help="Save only the first certificate as a PNG for checking layout")
    mode.add_argument("--send", action="store_true", help="Email PDFs using SENDER_EMAIL and SENDER_EMAIL_PASSWORD")
    parser.add_argument("--smtp-host", default="smtp.gmail.com", help="SMTP server with implicit TLS")
    parser.add_argument("--smtp-port", type=int, default=465, help="Implicit TLS SMTP port")
    parser.add_argument("--subject", default="Your certificate", help="Email subject")
    parser.add_argument("--message", default="Dear {name},\n\nPlease find your certificate attached.\nCertificate ID: {certificate_id}\n", help="Plain-text email body; supports {name} and {certificate_id}")
    parser.add_argument("--logo", type=Path, help="Optional inline email logo")
    return parser


def main(argv=None):
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        for option in ("start", "font_size", "detail_size", "name_width", "dpi", "smtp_port"):
            value = getattr(args, option)
            if value is not None and value < 1:
                raise ValueError(f"--{option.replace('_', '-')} must be positive.")
        if args.smtp_port > 65535:
            raise ValueError("--smtp-port must be between 1 and 65535.")
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,39}", args.prefix):
            raise ValueError("--prefix must start with a letter or digit and contain 1–40 letters, digits, underscores, or hyphens.")
        if any(ord(char) < 32 or ord(char) == 127 for char in args.date):
            raise ValueError("--date must be a single line without control characters.")
        if args.send:
            if not valid_email(os.environ.get("SENDER_EMAIL", "")) or not os.environ.get("SENDER_EMAIL_PASSWORD"):
                raise ValueError("Set SENDER_EMAIL and SENDER_EMAIL_PASSWORD before using --send.")
            if "\r" in args.subject or "\n" in args.subject:
                raise ValueError("--subject must be a single line.")
            try:
                for _, field, spec, conversion in Formatter().parse(args.message):
                    if field is not None and (field not in ("name", "certificate_id") or spec or conversion):
                        raise ValueError("Unsupported message placeholder")
            except ValueError as error:
                raise ValueError("--message supports only {name} and {certificate_id}; use {{ and }} for literal braces.") from error
            if args.logo:
                with Image.open(args.logo) as logo:
                    logo.verify()
        recipients = read_recipients(args.input, require_email=args.send)
        with Image.open(args.template) as source:
            template = source.convert("RGB")
        if args.preview:
            name, _ = recipients[0]
            image = render_certificate(template, name, f"{args.prefix}{args.start:03d}", args)
            with args.preview.open("xb") as destination:
                image.save(destination, format="PNG")
            print(f"Preview saved to {args.preview}")
            return 0

        args.output.mkdir(parents=True, exist_ok=True)
        if any(args.output.iterdir()):
            raise ValueError(f"Output directory {args.output} is not empty. Choose a new --output to avoid overwrites or duplicate email.")
        book = Workbook()
        sheet = book.active
        sheet.title = "Certificates"
        sheet.append(("Name", "Email", "Certificate ID", "PDF"))
        records = []
        for number, (name, email) in enumerate(recipients, start=args.start):
            certificate_id = f"{args.prefix}{number:03d}"
            safe_name = re.sub(r"[^\w-]+", "_", name).strip("_")
            safe_name = safe_name.encode("utf-8")[:120].decode("utf-8", errors="ignore") or "recipient"
            filename = f"certificate_{certificate_id}_{safe_name}.pdf"
            image = render_certificate(template, name, certificate_id, args)
            with (args.output / filename).open("xb") as destination:
                image.save(destination, format="PDF", resolution=args.dpi)
            records.append((name, email, certificate_id, filename))
            sheet.append(records[-1])
            # Preserve literal names such as '=Example' instead of creating formulas.
            for column in range(1, 5):
                sheet.cell(row=len(records) + 1, column=column).data_type = "s"
            print(f"[{number - args.start + 1}/{len(recipients)}] Created {filename}")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column, width in (("A", 35), ("B", 40), ("C", 25), ("D", 70)):
            sheet.column_dimensions[column].width = width
        with (args.output / "verify.xlsx").open("xb") as destination:
            book.save(destination)
        print(f"Created {len(records)} certificates and {args.output / 'verify.xlsx'}.")
        if args.send:
            failures = send_certificates(records, args)
            print(f"Email: {len(records) - failures} accepted, {failures} failed. See {args.output / 'delivery.csv'}.")
            return 1 if failures else 0
        return 0
    except (OSError, ValueError, BadZipFile, csv.Error, ParseError, DefusedXmlException, smtplib.SMTPException) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
