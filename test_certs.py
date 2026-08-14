"""Run with: python test_certs.py (temporary files and mocked SMTP only)."""

from contextlib import redirect_stderr, redirect_stdout
import csv
import io
import os
from pathlib import Path
import smtplib
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from PIL import Image, ImageChops

import certs


def main():
    with TemporaryDirectory() as directory:
        root = Path(directory)
        source = root / "people.csv"
        names = ["Ada McDonald", "../../O'Neil", "=Example", "A Very Long Recipient Name " * 5]
        emails = [f"Recipient{number}@Example.com" for number in range(len(names))]
        with source.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow((" Name ", "Email"))
            writer.writerows(zip(names, emails))
            writer.writerow(("", ""))
        template = root / "template.png"
        Image.new("RGBA", (1200, 800), "white").save(template)
        args = ["--input", str(source), "--template", str(template), "--prefix", "DEMO-", "--start", "98", "--detail-size", "18"]
        output = root / "certificates"
        with patch("certs.smtplib.SMTP_SSL") as smtp, redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            assert certs.main(args + ["--output", str(output)]) == 0
            smtp.assert_not_called()
            assert certs.main(args + ["--output", str(output)]) == 1, "Must not overwrite an existing batch"
        pdfs = sorted(output.glob("*.pdf"))
        assert len(pdfs) == len(names)
        assert all(pdf.read_bytes().startswith(b"%PDF-") for pdf in pdfs)
        assert all(pdf.parent == output for pdf in pdfs), "Names must not escape the output directory"
        book = load_workbook(output / "verify.xlsx")
        rows = list(book.active.values)[1:]
        assert [row[0] for row in rows] == [name.strip() for name in names]
        assert [row[1] for row in rows] == emails, "Email capitalization must be preserved"
        assert [row[2] for row in rows] == ["DEMO-098", "DEMO-099", "DEMO-100", "DEMO-101"]
        assert book.active["A4"].data_type == "s", "Names must not become spreadsheet formulas"
        book.close()
        assert certs.read_recipients(output / "verify.xlsx") == [(name.strip(), email) for name, email in zip(names, emails)]

        preview = root / "preview.png"
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            assert certs.main(args + ["--preview", str(preview)]) == 0
            original = preview.read_bytes()
            assert certs.main(args + ["--preview", str(preview)]) == 1
            assert preview.read_bytes() == original
        with Image.open(preview) as image:
            assert image.size == (1200, 800)
            band = image.crop((0, 200, 1200, 600))
            bbox = ImageChops.difference(band, Image.new("RGB", band.size, "white")).getbbox()
            assert bbox and abs((bbox[0] + bbox[2]) / 2 - 600) < 5, "Name must be centered"
            assert bbox[2] - bbox[0] <= 960, "Name must fit within 80% of the template"

        xlsx = root / "names.xlsx"
        book = Workbook()
        book.active.append(["Name"])
        book.active.append(["A Person"])
        book.save(xlsx)
        book.close()
        assert certs.read_recipients(xlsx) == [("A Person", "")]
        bad = root / "invalid.csv"
        for content in ("Email\na@example.com\n", "Name,Email\n,valid@example.com\n", "Name,Email\nPerson,bad-address\n", "Name,Name\nA,B\n", "Name,Email\n", 'Name,Email\n"Unclosed,person@example.com\n'):
            bad.write_text(content)
            invalid_output = root / "invalid-output"
            with redirect_stderr(io.StringIO()):
                assert certs.main(["--input", str(bad), "--output", str(invalid_output)]) == 1
            assert not invalid_output.exists(), "Invalid input must fail before writing anything"
        for extra in (["--start", "0"], ["--font-size", "0"], ["--prefix", "../oops"], ["--name-width", "-1"], ["--id-position", "9000", "9000"]):
            with redirect_stderr(io.StringIO()):
                assert certs.main(args + ["--preview", str(root / "invalid-preview.png")] + extra) == 1
            assert not (root / "invalid-preview.png").exists()

        sent_output = root / "sent"
        credentials = {"SENDER_EMAIL": "sender@example.com", "SENDER_EMAIL_PASSWORD": "test-password"}
        with patch.dict(os.environ, credentials), patch("certs.smtplib.SMTP_SSL") as smtp, redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            connection = smtp.return_value.__enter__.return_value
            connection.send_message.side_effect = [None, smtplib.SMTPException("test refusal"), None, None]
            assert certs.main(args + ["--output", str(sent_output), "--send", "--logo", str(template)]) == 1
            assert (sent_output / "verify.xlsx").exists(), "Generation must finish before delivery"
            assert connection.send_message.call_count == len(names), "One delivery failure must not stop the batch"
            message = connection.send_message.call_args_list[0].args[0]
            assert message["To"] == emails[0]
            assert "Ada McDonald" in message.get_body(preferencelist=("plain",)).get_content()
            attachments = list(message.iter_attachments())
            assert len(attachments) == 1 and attachments[0].get_content_type() == "application/pdf"
            assert attachments[0].get_payload(decode=True).startswith(b"%PDF-")
            assert any(part.get("Content-ID") == "<logo>" for part in message.walk())
        with (sent_output / "delivery.csv").open(newline="") as log:
            delivery = list(csv.DictReader(log))
        assert len(delivery) == len(names)
        assert delivery[1]["Status"].startswith("Failed")
        assert delivery[-1]["Status"] == "Accepted by SMTP server"

        with patch.dict(os.environ, credentials), patch("certs.smtplib.SMTP_SSL") as smtp, redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            assert certs.main(args + ["--output", str(root / "success"), "--send"]) == 0
            assert smtp.return_value.__enter__.return_value.send_message.call_count == len(names)
            smtp.reset_mock()
            for message in ("{name[foo]}", "{unknown}", "{name:100s}", "{name"):
                assert certs.main(args + ["--output", str(root / "bad-message"), "--send", "--message", message]) == 1
            smtp.assert_not_called()
            assert not (root / "bad-message").exists()

        with patch.dict(os.environ, {"SENDER_EMAIL": "", "SENDER_EMAIL_PASSWORD": ""}), patch("certs.smtplib.SMTP_SSL") as smtp, redirect_stderr(io.StringIO()):
            assert certs.main(args + ["--output", str(root / "missing-credentials"), "--send"]) == 1
            smtp.assert_not_called()
            assert not (root / "missing-credentials").exists()
        for email in ("", "@", "a@", "a@example.com\r\nBcc: b@example.com", "a@example.com,b@example.com", "not-an-address"):
            assert not certs.valid_email(email)
    print("Passed: generation, input validation, layout, IDs, overwrite protection, and mocked email delivery.")


if __name__ == "__main__":
    main()
